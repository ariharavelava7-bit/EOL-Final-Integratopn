"""
L&T CORe - Component Obsolescence & Resilience Engine
Main FastAPI Application with Authentication & Admin Panel
"""

from fastapi import FastAPI, HTTPException, Depends, status, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, EmailStr
from typing import List, Dict, Optional
from sqlalchemy.orm import Session
import os
from datetime import datetime, timedelta
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import google.generativeai as genai
from dotenv import load_dotenv

# Import database and auth
from database import (
    get_db, init_db, create_default_admin,
    User, UserStatus, UserRole, SearchHistory, SavedSearch, 
    Report, Notification, SystemNotification, AuditLog
)
from auth import (
    get_current_user, get_admin_user, authenticate_user, register_user,
    create_access_token, get_password_hash, verify_password,
    UserCreate, UserLogin, UserResponse, UserUpdate, PasswordChange, Token,
    ACCESS_TOKEN_EXPIRE_MINUTES
)

# Import API clients
from ari.api_clients import DigiKeyClient, OctopartClient

# Load environment variables
load_dotenv()

# Initialize FastAPI app
app = FastAPI(
    title="L&T CORe - Component Obsolescence & Resilience Engine",
    description="Complete part lifecycle management with authentication",
    version="2.0.0"
)

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configuration
OCTOPART_CLIENT_ID = os.getenv("OCTOPART_CLIENT_ID", "")
OCTOPART_CLIENT_SECRET = os.getenv("OCTOPART_CLIENT_SECRET", "")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
DIGIKEY_CLIENT_ID = os.getenv("DIGIKEY_CLIENT_ID", "")
DIGIKEY_CLIENT_SECRET = os.getenv("DIGIKEY_CLIENT_SECRET", "")

# Initialize external APIs
octopart_client = None
if OCTOPART_CLIENT_ID and OCTOPART_CLIENT_SECRET:
    octopart_client = OctopartClient(OCTOPART_CLIENT_ID, OCTOPART_CLIENT_SECRET)
    print("[OK] Octopart/Nexar API initialized (for comparison only)")

digikey_client = None
if DIGIKEY_CLIENT_ID and DIGIKEY_CLIENT_SECRET:
    digikey_client = DigiKeyClient(DIGIKEY_CLIENT_ID, DIGIKEY_CLIENT_SECRET)
    print("[OK] Digi-Key API initialized (for search + alternates)")

if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
    print("[OK] Gemini API initialized")


# ==================== Pydantic Models ====================

class PartSpec(BaseModel):
    parameter: str
    value: str

class CompareRequest(BaseModel):
    eol_part_number: str
    alternate_part_numbers: List[str]

class DownloadRequest(BaseModel):
    eol_part_number: str
    eol_specs: List[Dict]
    alternates: List[Dict]

class NotificationCreate(BaseModel):
    user_id: Optional[int] = None  # None = all users
    title: str
    message: str
    notification_type: str = "info"
    link: Optional[str] = None

class SystemNotificationCreate(BaseModel):
    title: str
    message: str
    notification_type: str = "info"
    expires_at: Optional[datetime] = None

class UserApproval(BaseModel):
    user_id: int
    action: str  # approve, reject, suspend

class SaveSearchRequest(BaseModel):
    name: str
    part_number: str
    notes: Optional[str] = None


# ==================== Startup Event ====================

@app.on_event("startup")
async def startup_event():
    """Initialize database and create default admin on startup"""
    print("\n" + "=" * 60)
    print("L&T CORe - Starting up...")
    print("=" * 60)
    init_db()
    create_default_admin()
    print("=" * 60 + "\n")


# ==================== Public Endpoints ====================

@app.get("/")
async def root():
    return {"message": "L&T CORe - Component Obsolescence & Resilience Engine API v2.0"}


@app.get("/api/health")
async def health_check():
    return {
        "status": "healthy",
        "version": "2.0.0",
        "digikey": "configured" if digikey_client else "not configured",
        "octopart": "configured" if octopart_client else "not configured"
    }


# ==================== Authentication Endpoints ====================

@app.post("/api/auth/register", response_model=dict)
async def register(user_data: UserCreate, db: Session = Depends(get_db)):
    """Register a new user (requires admin approval)"""
    try:
        new_user = register_user(db, user_data)
        
        # Create notification for admins
        admins = db.query(User).filter(User.role == UserRole.ADMIN.value).all()
        for admin in admins:
            notification = Notification(
                user_id=admin.id,
                title="New User Registration",
                message=f"{new_user.full_name} ({new_user.email}) has requested access.",
                notification_type="info",
                link="/admin/users"
            )
            db.add(notification)
        db.commit()
        
        return {
            "success": True,
            "message": "Registration successful! Please wait for admin approval.",
            "user_id": new_user.id
        }
    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/auth/login", response_model=Token)
async def login(credentials: UserLogin, db: Session = Depends(get_db)):
    """Login with email and password"""
    user = authenticate_user(db, credentials.email, credentials.password)
    
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect email or password"
        )
    
    if user.status == UserStatus.PENDING.value:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Your account is pending approval. Please wait for admin approval."
        )
    
    if user.status == UserStatus.REJECTED.value:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Your account request was rejected. Please contact administrator."
        )
    
    if user.status == UserStatus.SUSPENDED.value:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Your account has been suspended. Please contact administrator."
        )
    
    # Update last login
    user.last_login = datetime.utcnow()
    db.commit()
    
    # Create audit log
    audit = AuditLog(
        user_id=user.id,
        action="login",
        details=json.dumps({"email": user.email})
    )
    db.add(audit)
    db.commit()
    
    # Create access token
    access_token = create_access_token(
        data={"user_id": user.id, "email": user.email, "role": user.role}
    )
    
    return {
        "access_token": access_token,
        "token_type": "bearer",
        "user": {
            "id": user.id,
            "email": user.email,
            "username": user.username,
            "full_name": user.full_name,
            "role": user.role,
            "status": user.status,
            "department": user.department
        }
    }


@app.get("/api/auth/me", response_model=UserResponse)
async def get_current_user_info(current_user: User = Depends(get_current_user)):
    """Get current user information"""
    return current_user


@app.put("/api/auth/me", response_model=UserResponse)
async def update_profile(
    updates: UserUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Update current user profile"""
    if updates.full_name:
        current_user.full_name = updates.full_name
    if updates.department:
        current_user.department = updates.department
    if updates.phone:
        current_user.phone = updates.phone
    
    current_user.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(current_user)
    
    return current_user


@app.post("/api/auth/change-password")
async def change_password(
    password_data: PasswordChange,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Change password for current user"""
    if not verify_password(password_data.current_password, current_user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Current password is incorrect"
        )
    
    current_user.hashed_password = get_password_hash(password_data.new_password)
    current_user.updated_at = datetime.utcnow()
    db.commit()
    
    return {"success": True, "message": "Password changed successfully"}


# ==================== Admin Endpoints ====================

@app.get("/api/admin/users", response_model=List[dict])
async def get_all_users(
    status_filter: Optional[str] = None,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Get all users (admin only)"""
    query = db.query(User)
    
    if status_filter:
        query = query.filter(User.status == status_filter)
    
    users = query.order_by(User.created_at.desc()).all()
    
    return [
        {
            "id": u.id,
            "email": u.email,
            "username": u.username,
            "full_name": u.full_name,
            "role": u.role,
            "status": u.status,
            "department": u.department,
            "phone": u.phone,
            "created_at": u.created_at.isoformat() if u.created_at else None,
            "last_login": u.last_login.isoformat() if u.last_login else None,
            "approved_at": u.approved_at.isoformat() if u.approved_at else None
        }
        for u in users
    ]


@app.get("/api/admin/pending-users", response_model=List[dict])
async def get_pending_users(
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Get all pending user registrations (admin only)"""
    users = db.query(User).filter(User.status == UserStatus.PENDING.value).all()
    
    return [
        {
            "id": u.id,
            "email": u.email,
            "username": u.username,
            "full_name": u.full_name,
            "department": u.department,
            "phone": u.phone,
            "created_at": u.created_at.isoformat() if u.created_at else None
        }
        for u in users
    ]


@app.post("/api/admin/users/approve")
async def approve_user(
    approval: UserApproval,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Approve, reject, or suspend a user (admin only)"""
    user = db.query(User).filter(User.id == approval.user_id).first()
    
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    if approval.action == "approve":
        user.status = UserStatus.APPROVED.value
        user.approved_at = datetime.utcnow()
        user.approved_by = admin.id
        message = "Your account has been approved! You can now log in."
        notif_type = "success"
    elif approval.action == "reject":
        user.status = UserStatus.REJECTED.value
        message = "Your account request has been rejected."
        notif_type = "warning"
    elif approval.action == "suspend":
        user.status = UserStatus.SUSPENDED.value
        message = "Your account has been suspended."
        notif_type = "warning"
    else:
        raise HTTPException(status_code=400, detail="Invalid action")
    
    # Create notification for user
    notification = Notification(
        user_id=user.id,
        title="Account Status Update",
        message=message,
        notification_type=notif_type
    )
    db.add(notification)
    
    # Audit log
    audit = AuditLog(
        user_id=admin.id,
        action=f"user_{approval.action}",
        details=json.dumps({"target_user_id": user.id, "target_email": user.email})
    )
    db.add(audit)
    
    db.commit()
    
    return {"success": True, "message": f"User {approval.action}d successfully"}


@app.put("/api/admin/users/{user_id}/role")
async def update_user_role(
    user_id: int,
    role: str,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Update user role (admin only)"""
    if role not in [r.value for r in UserRole]:
        raise HTTPException(status_code=400, detail="Invalid role")
    
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    user.role = role
    user.updated_at = datetime.utcnow()
    db.commit()
    
    return {"success": True, "message": f"User role updated to {role}"}


@app.delete("/api/admin/users/{user_id}")
async def delete_user(
    user_id: int,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Delete a user (admin only)"""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    if user.role == UserRole.ADMIN.value:
        # Check if this is the last admin
        admin_count = db.query(User).filter(User.role == UserRole.ADMIN.value).count()
        if admin_count <= 1:
            raise HTTPException(status_code=400, detail="Cannot delete the last admin")
    
    db.delete(user)
    db.commit()
    
    return {"success": True, "message": "User deleted successfully"}


# ==================== Notification Endpoints ====================

@app.get("/api/notifications")
async def get_user_notifications(
    unread_only: bool = False,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get notifications for current user"""
    query = db.query(Notification).filter(Notification.user_id == current_user.id)
    
    if unread_only:
        query = query.filter(Notification.is_read == False)
    
    notifications = query.order_by(Notification.created_at.desc()).limit(50).all()
    
    return [
        {
            "id": n.id,
            "title": n.title,
            "message": n.message,
            "type": n.notification_type,
            "is_read": n.is_read,
            "link": n.link,
            "created_at": n.created_at.isoformat()
        }
        for n in notifications
    ]


@app.put("/api/notifications/{notification_id}/read")
async def mark_notification_read(
    notification_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Mark a notification as read"""
    notification = db.query(Notification).filter(
        Notification.id == notification_id,
        Notification.user_id == current_user.id
    ).first()
    
    if not notification:
        raise HTTPException(status_code=404, detail="Notification not found")
    
    notification.is_read = True
    notification.read_at = datetime.utcnow()
    db.commit()
    
    return {"success": True}


@app.put("/api/notifications/read-all")
async def mark_all_notifications_read(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Mark all notifications as read"""
    db.query(Notification).filter(
        Notification.user_id == current_user.id,
        Notification.is_read == False
    ).update({"is_read": True, "read_at": datetime.utcnow()})
    db.commit()
    
    return {"success": True}


@app.post("/api/admin/notifications/send")
async def send_notification(
    notification: NotificationCreate,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Send notification to user(s) (admin only)"""
    if notification.user_id:
        # Send to specific user
        user = db.query(User).filter(User.id == notification.user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        
        notif = Notification(
            user_id=notification.user_id,
            title=notification.title,
            message=notification.message,
            notification_type=notification.notification_type,
            link=notification.link
        )
        db.add(notif)
    else:
        # Send to all approved users
        users = db.query(User).filter(User.status == UserStatus.APPROVED.value).all()
        for user in users:
            notif = Notification(
                user_id=user.id,
                title=notification.title,
                message=notification.message,
                notification_type=notification.notification_type,
                link=notification.link
            )
            db.add(notif)
    
    db.commit()
    return {"success": True, "message": "Notification sent successfully"}


@app.post("/api/admin/system-notifications")
async def create_system_notification(
    notification: SystemNotificationCreate,
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Create a system-wide notification (admin only)"""
    sys_notif = SystemNotification(
        title=notification.title,
        message=notification.message,
        notification_type=notification.notification_type,
        created_by=admin.id,
        expires_at=notification.expires_at
    )
    db.add(sys_notif)
    db.commit()
    
    return {"success": True, "message": "System notification created"}


@app.get("/api/system-notifications")
async def get_system_notifications(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get active system notifications"""
    now = datetime.utcnow()
    notifications = db.query(SystemNotification).filter(
        SystemNotification.is_active == True,
        (SystemNotification.expires_at == None) | (SystemNotification.expires_at > now)
    ).order_by(SystemNotification.created_at.desc()).all()
    
    return [
        {
            "id": n.id,
            "title": n.title,
            "message": n.message,
            "type": n.notification_type,
            "created_at": n.created_at.isoformat()
        }
        for n in notifications
    ]


# ==================== Search History & Saved Searches ====================

@app.get("/api/search-history")
async def get_search_history(
    limit: int = 20,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get user's search history"""
    history = db.query(SearchHistory).filter(
        SearchHistory.user_id == current_user.id
    ).order_by(SearchHistory.searched_at.desc()).limit(limit).all()
    
    return [
        {
            "id": h.id,
            "part_number": h.part_number,
            "manufacturer": h.manufacturer,
            "description": h.description,
            "alternates_count": h.alternates_count,
            "searched_at": h.searched_at.isoformat()
        }
        for h in history
    ]


@app.post("/api/saved-searches")
async def save_search(
    search_data: SaveSearchRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Save a search for quick access"""
    saved = SavedSearch(
        user_id=current_user.id,
        name=search_data.name,
        part_number=search_data.part_number,
        notes=search_data.notes
    )
    db.add(saved)
    db.commit()
    
    return {"success": True, "message": "Search saved successfully", "id": saved.id}


@app.get("/api/saved-searches")
async def get_saved_searches(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get user's saved searches"""
    saved = db.query(SavedSearch).filter(
        SavedSearch.user_id == current_user.id
    ).order_by(SavedSearch.created_at.desc()).all()
    
    return [
        {
            "id": s.id,
            "name": s.name,
            "part_number": s.part_number,
            "notes": s.notes,
            "created_at": s.created_at.isoformat()
        }
        for s in saved
    ]


@app.delete("/api/saved-searches/{search_id}")
async def delete_saved_search(
    search_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Delete a saved search"""
    saved = db.query(SavedSearch).filter(
        SavedSearch.id == search_id,
        SavedSearch.user_id == current_user.id
    ).first()
    
    if not saved:
        raise HTTPException(status_code=404, detail="Saved search not found")
    
    db.delete(saved)
    db.commit()
    
    return {"success": True}


# ==================== Reports ====================

@app.get("/api/reports")
async def get_user_reports(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get user's generated reports"""
    reports = db.query(Report).filter(
        Report.user_id == current_user.id
    ).order_by(Report.created_at.desc()).limit(50).all()
    
    return [
        {
            "id": r.id,
            "report_name": r.report_name,
            "report_type": r.report_type,
            "part_number": r.part_number,
            "created_at": r.created_at.isoformat()
        }
        for r in reports
    ]


# ==================== Admin Dashboard Stats ====================

@app.get("/api/admin/stats")
async def get_admin_stats(
    admin: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    """Get admin dashboard statistics"""
    total_users = db.query(User).count()
    pending_users = db.query(User).filter(User.status == UserStatus.PENDING.value).count()
    approved_users = db.query(User).filter(User.status == UserStatus.APPROVED.value).count()
    
    # Active users (logged in within last 7 days)
    week_ago = datetime.utcnow() - timedelta(days=7)
    active_users = db.query(User).filter(
        User.last_login >= week_ago,
        User.status == UserStatus.APPROVED.value
    ).count()
    
    total_searches = db.query(SearchHistory).count()
    total_reports = db.query(Report).count()
    
    # Recent activity
    recent_searches = db.query(SearchHistory).order_by(
        SearchHistory.searched_at.desc()
    ).limit(10).all()
    
    return {
        "users": {
            "total": total_users,
            "pending": pending_users,
            "approved": approved_users,
            "active_this_week": active_users
        },
        "activity": {
            "total_searches": total_searches,
            "total_reports": total_reports
        },
        "recent_searches": [
            {
                "part_number": s.part_number,
                "user_id": s.user_id,
                "searched_at": s.searched_at.isoformat()
            }
            for s in recent_searches
        ]
    }


# ==================== Part Lookup Endpoints (Protected) ====================

@app.get("/api/v1/lookup_eol_specs/{part_number}")
async def lookup_eol_specs(
    part_number: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Main endpoint for part lookup - ONLY uses Digi-Key
    Requires authentication
    """
    try:
        print(f"\n{'='*60}")
        print(f"[Search] User: {current_user.email} | Part: {part_number}")
        print(f"{'='*60}")
        
        if not digikey_client:
            raise HTTPException(
                status_code=503, 
                detail="Digi-Key API not configured"
            )
        
        part_data = {
            "partNumber": part_number,
            "manufacturer": "Unknown",
            "description": "Unknown",
            "status": "Active",
            "lastChecked": datetime.now().strftime("%b %d, %Y"),
            "specs": [],
            "lifecycle": {
                "status": "Active",
                "yearsToEOL": "N/A",
                "introduced": "N/A",
                "source": "Digi-Key"
            },
            "compliance": {
                "rohs": {"status": "Unknown", "directive": ""},
                "reachSvhc": "Unknown",
                "tscaPbt": "Unknown",
                "globalPfas": "Unknown"
            },
            "marketAvailability": {
                "distributors": 0,
                "availQty": "0",
                "minPrice": "N/A",
                "minLeadtime": "N/A"
            },
            "alternates": []
        }
        
        # Step 1: Get product details from Digi-Key
        print("\n[Step 1] Fetching product details from Digi-Key...")
        dk_product = digikey_client.search_part(part_number)
        
        if dk_product:
            part_data["manufacturer"] = dk_product.get("Manufacturer", "Unknown")
            part_data["description"] = dk_product.get("Description", "Unknown")
            
            for key, value in dk_product.items():
                if key not in ["_source", "_cached"] and value and value != "N/A":
                    part_data["specs"].append({
                        "parameter": key.replace("SPEC_", ""),
                        "value": str(value)
                    })
            
            if dk_product.get("QuantityAvailable"):
                part_data["marketAvailability"]["availQty"] = str(dk_product.get("QuantityAvailable", 0))
                part_data["marketAvailability"]["distributors"] = 1
            
            if dk_product.get("UnitPrice"):
                part_data["marketAvailability"]["minPrice"] = f"${dk_product.get('UnitPrice')}"
        
        # Step 2: Get alternate packagings from Digi-Key
        print("\n[Step 2] Fetching alternate packagings from Digi-Key...")
        alternates = digikey_client.get_alternate_packaging(part_number, limit=5)
        
        for alt in alternates:
            part_data["alternates"].append({
                "partNumber": alt.get("ManufacturerPartNumber", ""),
                "manufacturer": alt.get("Manufacturer", ""),
                "description": alt.get("Description", ""),
                "unitPrice": alt.get("UnitPrice", "N/A"),
                "quantityAvailable": alt.get("QuantityAvailable", 0),
                "digiKeyPartNumber": alt.get("DigiKeyProductNumber", ""),
                "productUrl": alt.get("ProductUrl", ""),
                "crossType": "SF",
                "acl": "No",
                "yearsToEOL": "N/A",
                "inventory": str(alt.get("QuantityAvailable", "--")),
                "difference": "--"
            })
        
        # Save to search history
        history = SearchHistory(
            user_id=current_user.id,
            part_number=part_number,
            manufacturer=part_data["manufacturer"],
            description=part_data["description"][:500] if part_data["description"] else None,
            alternates_count=len(part_data["alternates"])
        )
        db.add(history)
        db.commit()
        
        if not part_data["specs"] and not part_data["alternates"]:
            raise HTTPException(status_code=404, detail=f"No data found for {part_number}")
        
        return part_data
    
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/v1/compare_parts")
async def compare_parts(
    request: CompareRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Compare EOL part with alternates using Octopart"""
    try:
        print(f"\n[Compare] User: {current_user.email}")
        print(f"EOL Part: {request.eol_part_number}")
        
        comparison_data = {
            "eolPart": None,
            "alternates": [],
            "specs": []
        }
        
        if not octopart_client:
            return comparison_data
        
        # Get EOL part specs from Octopart
        eol_parts = octopart_client.search_parts(request.eol_part_number, limit=1)
        if eol_parts:
            comparison_data["eolPart"] = eol_parts[0]
        
        # Get each alternate's specs from Octopart
        for alt_pn in request.alternate_part_numbers:
            alt_parts = octopart_client.search_parts(alt_pn, limit=1)
            if alt_parts:
                comparison_data["alternates"].append(alt_parts[0])
            else:
                comparison_data["alternates"].append({
                    "ManufacturerPartNumber": alt_pn,
                    "Manufacturer": "Unknown",
                    "Description": "Details not available"
                })
        
        # Build comparison specs
        all_specs = set()
        if comparison_data["eolPart"]:
            for key in comparison_data["eolPart"].keys():
                if not key.startswith("_"):
                    all_specs.add(key)
        
        for alt in comparison_data["alternates"]:
            for key in alt.keys():
                if not key.startswith("_"):
                    all_specs.add(key)
        
        for spec in sorted(all_specs):
            row = {
                "parameter": spec,
                "eolValue": str(comparison_data["eolPart"].get(spec, "N/A")) if comparison_data["eolPart"] else "N/A",
                "alternateValues": []
            }
            for alt in comparison_data["alternates"]:
                row["alternateValues"].append(str(alt.get(spec, "N/A")))
            comparison_data["specs"].append(row)
        
        return comparison_data
    
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/v1/download_report")
async def download_report(
    request: DownloadRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate and download Excel report"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Part Comparison"
        
        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        match_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        diff_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:E1')
        ws['A1'] = f"L&T CORe - Part Comparison: {request.eol_part_number}"
        ws['A1'].font = Font(bold=True, size=14)
        
        # Headers
        headers = ["Parameter", f"EOL: {request.eol_part_number}"]
        for i, alt in enumerate(request.alternates[:4]):
            headers.append(f"Alt {i+1}: {alt.get('partNumber', 'N/A')}")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        
        # Data rows
        row_num = 4
        eol_spec_dict = {spec.get("parameter"): spec.get("value") for spec in request.eol_specs}
        
        for param in sorted(set(s.get("parameter", "") for s in request.eol_specs)):
            if not param:
                continue
            
            ws.cell(row=row_num, column=1, value=param).border = thin_border
            eol_value = eol_spec_dict.get(param, "N/A")
            ws.cell(row=row_num, column=2, value=str(eol_value)).border = thin_border
            
            for col, alt in enumerate(request.alternates[:4], 3):
                alt_value = alt.get(param, "N/A")
                cell = ws.cell(row=row_num, column=col, value=str(alt_value))
                cell.border = thin_border
                if str(alt_value) == str(eol_value) and eol_value != "N/A":
                    cell.fill = match_fill
                elif alt_value != "N/A":
                    cell.fill = diff_fill
            
            row_num += 1
        
        # Column widths
        ws.column_dimensions['A'].width = 30
        for col in range(2, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 25
        
        # Save file
        os.makedirs("reports", exist_ok=True)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"EOL_Comparison_{request.eol_part_number}_{timestamp}.xlsx"
        filepath = os.path.join("reports", filename)
        wb.save(filepath)
        
        # Save to database
        report = Report(
            user_id=current_user.id,
            report_name=filename,
            report_type="comparison",
            file_path=filepath,
            part_number=request.eol_part_number,
            alternates_included=json.dumps([a.get("partNumber") for a in request.alternates[:4]])
        )
        db.add(report)
        db.commit()
        
        return FileResponse(
            path=filepath,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=filename
        )
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


# ==================== Main Entry ====================

if __name__ == "__main__":
    import uvicorn
    print("=" * 60)
    print("L&T CORe v2.0 - With Authentication & Admin Panel")
    print("=" * 60)
    print(f"Digi-Key:  {'[OK]' if digikey_client else '[NOT CONFIGURED]'}")
    print(f"Octopart:  {'[OK]' if octopart_client else '[NOT CONFIGURED]'}")
    print("=" * 60)
    print("Default Admin: admin@ltcore.com / admin123")
    print("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=8001)
